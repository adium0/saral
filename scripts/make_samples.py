import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from sample_text import CELL, OS, CPU
from reportlab.lib.pagesizes import A4
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.units import mm

OUT = os.path.join(os.path.dirname(__file__), '..', 'public', 'samples')
os.makedirs(os.path.join(OUT, 'test-cases'), exist_ok=True)

body = ParagraphStyle('b', fontName='Times-Roman', fontSize=10.5, leading=12.2, alignment=TA_JUSTIFY, firstLineIndent=12, spaceAfter=2)
h1 = ParagraphStyle('h1', fontName='Times-Bold', fontSize=18, leading=22, spaceAfter=8)
h2 = ParagraphStyle('h2', fontName='Times-Bold', fontSize=12.5, leading=15, spaceBefore=8, spaceAfter=4)

def make_pdf(spec, path, cols=1, margin=34*mm):
    def deco(c, d):
        c.saveState()
        c.setFont('Times-Italic', 8.5)
        c.drawString(margin, A4[1] - 12*mm, spec['header'])
        c.drawRightString(A4[0] - margin, 10*mm, str(d.page))
        c.restoreState()
    doc = BaseDocTemplate(path, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=20*mm, bottomMargin=18*mm, title=spec['title'], author='Saral team')
    w = (A4[0] - 2*margin)
    if cols == 1:
        frames = [Frame(margin, 18*mm, w, A4[1]-38*mm, id='f')]
    else:
        gap = 6*mm; cw = (w - gap)/2
        frames = [Frame(margin, 18*mm, cw, A4[1]-38*mm, id='a'), Frame(margin+cw+gap, 18*mm, cw, A4[1]-38*mm, id='b')]
    doc.addPageTemplates([PageTemplate(id='p', frames=frames, onPage=deco)])
    story = [Paragraph(spec['title'], h1)]
    for heading, paras in spec['sections']:
        if heading != spec['title']:
            story.append(Paragraph(heading, h2))
        for p in paras:
            story.append(Paragraph(p, body))
    doc.build(story)

make_pdf(CELL, os.path.join(OUT, 'cell-structure.pdf'))
make_pdf(OS, os.path.join(OUT, 'operating-systems.pdf'))
make_pdf(OS, os.path.join(OUT, 'test-cases', 'two-column.pdf'), cols=2, margin=15*mm)

# DOCX lecture notes with headings and lists
from docx import Document
d = Document()
d.core_properties.title = CPU['title']
for i, (heading, items) in enumerate(CPU['sections']):
    d.add_heading(heading, level=1 if i == 0 else 2)
    for kind, text in items:
        d.add_paragraph(text, style='List Bullet' if kind == 'li' else None)
d.save(os.path.join(OUT, 'cpu-scheduling.docx'))

# "Scanned" PDF: the page is only an image, with no text layer.
from PIL import Image, ImageDraw, ImageFont
img = Image.new('L', (1240, 1754), 245)
dr = ImageDraw.Draw(img)
try:
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf', 26)
except Exception:
    font = ImageFont.load_default()
y = 120
import textwrap
for heading, paras in CELL['sections'][:2]:
    dr.text((110, y), heading, fill=20, font=font); y += 60
    for p in paras:
        for line in textwrap.wrap(p, 78):
            dr.text((110, y), line, fill=40, font=font); y += 36
        y += 18
img = img.rotate(0.8, fillcolor=245)
img.save(os.path.join(OUT, 'test-cases', 'scanned-page.pdf'), 'PDF', resolution=150)
print('done')
